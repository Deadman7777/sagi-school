"""État des services optionnels : qui les prend, ce qu'ils doivent et paient.

On teste la COHÉRENCE avec les autres écrans (échéancier, guichet) plutôt
que des valeurs isolées : un état qui contredit la fiche ne sert à rien.
"""
import datetime
from io import BytesIO

from rest_framework.test import APITestCase

from apps.eleves.etat_services import etat_services
from apps.eleves.models import Eleve, EleveService, Section, Service
from apps.paiements.models import Exercice, Receveur
from apps.tenants.models import Tenant
from apps.users.models import User

RENTREE = datetime.date(2026, 10, 1)
AUJOURDHUI = datetime.date(2026, 11, 15)            # octobre et novembre échus


class EtatServicesTest(APITestCase):
    def setUp(self):
        self.tenant = Tenant.objects.create(nom='Les Pins', code_etablissement='PIN')
        user = User.objects.create_user('a@pins.sn', 'x', nom='Caissière', role='ADMIN_ECOLE',
                                        tenant=self.tenant)
        self.client.force_authenticate(user)
        self.ex = Exercice.objects.create(tenant=self.tenant, annee_scolaire='2026-2027',
                                          nb_mensualites=9, date_debut=RENTREE,
                                          date_fin=datetime.date(2027, 6, 30))
        self.cp = Section.objects.create(tenant=self.tenant, nom='CP', frais_mensualite=20000)
        self.transport = Service.objects.create(tenant=self.tenant, nom='Transport',
                                                montant=30000, periodicite='MENSUEL')
        self.cantine = Service.objects.create(tenant=self.tenant, nom='Cantine',
                                              montant=10000, periodicite='MENSUEL')
        self.awa = self._eleve('Awa NDIAYE', self.transport, self.cantine)
        self.binta = self._eleve('Binta FALL', self.transport)
        self.coumba = self._eleve('Coumba SOW')                      # aucun service
        self.laurence = Receveur.objects.create(tenant=self.tenant, nom='Laurence')

    def _eleve(self, nom, *services):
        e = Eleve.objects.create(tenant=self.tenant, exercice=self.ex, section=self.cp,
                                 nom_complet=nom, date_inscription=RENTREE)
        for s in services:
            EleveService.objects.create(tenant=self.tenant, eleve=e, service=s)
        return e

    def _transport(self, eleve, montant, mois, **corps):
        r = self.client.post('/api/paiements/paiements/', {
            'eleve': str(eleve.id), 'mode_paiement': corps.pop('mode', 'ESPECE'),
            'montant_divers': montant, 'mois_regles': mois,
            'services_regles': [{'nom': 'Transport', 'montant': montant, 'nature': 'MENSUEL',
                                 'service': str(self.transport.id)}], **corps}, format='json')
        self.assertEqual(r.status_code, 201, r.content[:300])

    def _etat(self, **kw):
        return etat_services(self.tenant, self.ex, today=AUJOURDHUI, **kw)

    def _service(self, etat, nom):
        return next(s for s in etat['services'] if s['nom'] == nom)

    def test_abonnes_par_service(self):
        etat = self._etat()
        self.assertEqual(self._service(etat, 'Transport')['nb_abonnes'], 2)
        self.assertEqual(self._service(etat, 'Cantine')['nb_abonnes'], 1)
        noms = [a['nom_complet'] for a in self._service(etat, 'Transport')['abonnes']]
        self.assertEqual(sorted(noms), ['Awa NDIAYE', 'Binta FALL'])
        self.assertEqual(self._service(etat, 'Transport')['abonnes'][0]['section'], 'CP')

    def test_du_coherent_avec_le_total_annuel_de_la_fiche(self):
        etat = self._etat()
        awa = Eleve.objects.get(pk=self.awa.pk)
        du_services = sum(a['du_annee'] for s in etat['services'] for a in s['abonnes']
                          if a['id'] == str(awa.id))
        self.assertEqual(du_services, awa.montant_services_annuel)
        # Deux mois échus (octobre, novembre) au 15 novembre.
        tr = next(a for a in self._service(etat, 'Transport')['abonnes'] if a['id'] == str(awa.id))
        self.assertEqual((tr['du_echu'], tr['nb_mois']), (60000, 9))

    def test_paye_reste_et_tresorerie_par_receveur(self):
        self._transport(self.awa, 30000, [10], receveur=str(self.laurence.id))
        self._transport(self.binta, 10000, [10], mode='WAVE')
        tr = self._service(self._etat(), 'Transport')
        awa = next(a for a in tr['abonnes'] if a['nom_complet'] == 'Awa NDIAYE')
        self.assertEqual((awa['paye'], awa['reste_echu'], awa['statut']), (30000, 30000, 'PARTIEL'))
        self.assertEqual(tr['paye'], 40000)
        self.assertEqual(tr['encaisse'], 40000)
        self.assertEqual(sum(m['montant'] for m in tr['par_mode']), tr['encaisse'])
        self.assertEqual({r['nom']: r['montant'] for r in tr['par_receveur']},
                         {'Laurence': 30000, 'Caissière': 10000})
        self.assertEqual(tr['reste_echu'], tr['du_echu'] - tr['paye'])

    def test_periode_des_encaissements(self):
        self._transport(self.awa, 30000, [10], date_paiement='2026-10-05')
        self._transport(self.binta, 30000, [10], date_paiement='2026-11-05')
        tr = self._service(self._etat(du=datetime.date(2026, 11, 1)), 'Transport')
        self.assertEqual(tr['encaisse'], 30000)
        self.assertEqual(tr['paye'], 60000)         # le payé de l'année n'est pas borné

    def test_filtre_par_service_et_exports(self):
        self._transport(self.awa, 30000, [10])
        r = self.client.get('/api/eleves/etat-services/', {'service': str(self.transport.id)})
        self.assertEqual([s['nom'] for s in r.json()['services']], ['Transport'])
        pdf = self.client.get('/api/eleves/etat-services/', {'export': 'pdf'})
        self.assertEqual(pdf.status_code, 200)
        from pypdf import PdfReader
        texte = ''.join(p.extract_text() for p in PdfReader(BytesIO(pdf.content)).pages)
        self.assertIn('Awa NDIAYE', texte)
        self.assertIn('Transport', texte)
        self.assertNotIn('{#', texte)
        xlsx = self.client.get('/api/eleves/etat-services/', {'export': 'xlsx'})
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(xlsx.content))
        self.assertEqual(wb.sheetnames, ['Récapitulatif', 'Cantine', 'Transport'])
