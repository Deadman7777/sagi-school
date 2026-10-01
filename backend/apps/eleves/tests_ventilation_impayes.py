"""Tests : ce qui reste dû, en trois parts, sur les documents remis aux familles.

Demandé le 01/10/2026 : avant le total, la fiche de situation (élève comme
famille) doit séparer les impayés des années antérieures, les impayés de
l'année déjà échus (mois en retard ou partiels) et ce qui n'est pas encore
échu — pour qu'un parent comprenne ce qu'on lui réclame.
"""
import datetime
from io import BytesIO

from pypdf import PdfReader
from rest_framework.test import APITestCase

from apps.eleves.echeancier import construire_echeancier, ventilation_du
from apps.eleves.models import Eleve, Famille, Section
from apps.paiements.models import Exercice, Paiement
from apps.tenants.models import Tenant
from apps.users.models import User

MARS = datetime.date(2026, 3, 15)


class _Ecole(APITestCase):
    """Le décor commun : une école, un exercice, une section, une famille."""

    def setUp(self):
        self.tenant = Tenant.objects.create(nom='LPE', code_etablissement='LPE')
        self.user = User.objects.create_user('d@lpe.sn', 'x', nom='Dir', role='ADMIN_ECOLE',
                                             tenant=self.tenant)
        self.client.force_authenticate(self.user)
        self.ex = Exercice.objects.create(
            tenant=self.tenant, annee_scolaire='2025-2026', nb_mensualites=10,
            date_debut=datetime.date(2025, 10, 1), date_fin=datetime.date(2026, 7, 31))
        self.section = Section.objects.create(tenant=self.tenant, nom='CP', frais_inscription=20000,
                                              frais_mensualite=10000)
        self.famille = Famille.objects.create(tenant=self.tenant, nom='Famille DIOUF')

    def _eleve(self, nom, **extra):
        return Eleve.objects.create(tenant=self.tenant, exercice=self.ex, section=self.section,
                                    nom_complet=nom, famille=self.famille,
                                    date_inscription=self.ex.date_debut, **extra)

    def _payer(self, eleve, **montants):
        return Paiement.objects.create(tenant=self.tenant, exercice=self.ex, eleve=eleve,
                                       no_piece=f'REC-{Paiement.objects.count() + 1:04d}',
                                       **montants)


class VentilationTest(_Ecole):
    def test_les_trois_parts_font_le_total(self):
        e = self._eleve('Awa DIOUF', reliquat_anterieur=100000)
        self._payer(e, montant_inscription=20000, montant_mensualite=25000)  # 2,5 mois

        ech = construire_echeancier(e, today=MARS)
        v = ventilation_du(e, ech)

        self.assertEqual(v['anterieur'], 100000)
        self.assertGreater(v['annee_echue'], 0)
        self.assertGreater(v['a_venir'], 0)
        # L'année : ce qui reste sur l'échéancier, ni plus ni moins.
        self.assertEqual(v['annee_echue'] + v['a_venir'], ech['totaux']['reste'])
        self.assertEqual(v['exigible'], v['anterieur'] + v['annee_echue'])
        self.assertEqual(v['total'], v['anterieur'] + v['annee_echue'] + v['a_venir'])
        # Le même total que la fiche élève (un seul calcul).
        self.assertEqual(v['total'], ech['synthese']['total_restant_du_famille'])

    def test_une_fiche_de_creance_garde_son_ardoise(self):
        """L'enfant parti avec une dette : pas d'échéancier, mais l'ardoise est réelle."""
        e = self._eleve('Moussa DIOUF', reliquat_anterieur=60000, fiche_creance=True,
                        statut='TRANSFERE')
        v = ventilation_du(e)
        self.assertEqual((v['anterieur'], v['annee_echue'], v['a_venir'], v['total']),
                         (60000, 0, 0, 60000))

    def _texte_pdf(self, url):
        r = self.client.get(url)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return ''.join(p.extract_text() for p in PdfReader(BytesIO(r.content)).pages)

    def test_la_fiche_eleve_decompose_avant_le_total(self):
        e = self._eleve('Awa DIOUF', reliquat_anterieur=100000)
        texte = self._texte_pdf(f'/api/eleves/{e.id}/situation-pdf/')
        self.assertIn('Ce qui reste dû', texte)
        self.assertIn('Impayés des années antérieures', texte)
        self.assertIn("Impayés de l'année en cours", texte)
        self.assertIn('Non encore échu', texte)
        self.assertLess(texte.index('Ce qui reste dû'), texte.index('TOTAL RESTANT DÛ'))
        self.assertNotIn('{%', texte)

    def test_la_fiche_eleve_ignore_les_recus_annules(self):
        e = self._eleve('Awa DIOUF')
        p = self._payer(e, montant_inscription=20000)
        p.statut = 'ANNULE'
        p.save()
        texte = self._texte_pdf(f'/api/eleves/{e.id}/situation-pdf/')
        self.assertNotIn(p.no_piece, texte)

    def test_la_fiche_famille_decompose_par_enfant(self):
        self._eleve('Awa DIOUF', reliquat_anterieur=100000)
        self._eleve('Moussa DIOUF', reliquat_anterieur=60000, fiche_creance=True,
                    statut='TRANSFERE')
        texte = self._texte_pdf(f'/api/eleves/familles/{self.famille.id}/situation-pdf/')
        self.assertIn('Ce qui reste dû', texte)
        self.assertIn('Total famille', texte)
        # L'ardoise du créancier compte dans les antérieurs de la famille.
        self.assertIn('160 000', texte)


class EtatImpayesTest(_Ecole):
    """L'état du comité : chaque statut, sa somme, les sous-totaux, dans l'ordre
    de recouvrement — au 15 mars, mensualité exigible en début de mois."""

    def _paye(self, nom, mois_payes, reliquat=0):
        e = self._eleve(nom, reliquat_anterieur=reliquat)
        self._payer(e, montant_inscription=20000, montant_mensualite=10000 * mois_payes)
        return e

    def _etat(self):
        from apps.eleves.etat_impayes import etat_impayes
        return etat_impayes(self.tenant, self.ex, today=MARS)

    def _ligne(self, etat, code):
        return next(l for l in etat['lignes'] if l['code'] == code)

    def setUp(self):
        super().setUp()
        self.critique = self._eleve('Aliou CRITIQUE')                  # rien payé
        self.urgent = self._paye('Binta URGENT', 4)                    # oct → janv : fév + mars dus
        self.mois = self._paye('Coumba MOISENCOURS', 5)                # oct → fév : mars seul
        self.retard = self._paye('Daouda ARDOISE', 6, reliquat=50000)  # à jour, ardoise ancienne
        self.ok = self._paye('Fatou AJOUR', 6)                         # à jour

    def test_chaque_eleve_est_dans_son_groupe(self):
        etat = self._etat()
        noms = {l['code']: [e['nom_complet'] for e in l.get('eleves', [])] for l in etat['lignes']}
        self.assertEqual(noms['CRITIQUE'], ['Aliou CRITIQUE'])
        self.assertEqual(noms['URGENT'], ['Binta URGENT'])
        self.assertEqual(noms['ATTENTION_MOIS'], ['Coumba MOISENCOURS'])
        self.assertEqual(noms['ATTENTION_RETARD'], ['Daouda ARDOISE'])
        self.assertEqual((etat['nb_eleves'], etat['nb_a_jour']), (5, 1))

    def test_l_ordre_de_recouvrement_et_les_sous_totaux(self):
        etat = self._etat()
        self.assertEqual([l['code'] for l in etat['lignes']],
                         ['CRITIQUE', 'URGENT', 'PRIORITAIRES',
                          'ATTENTION_RETARD', 'ATTENTION_MOIS', 'ATTENTION'])
        prio = self._ligne(etat, 'PRIORITAIRES')
        self.assertEqual(prio['nb'], 2)
        self.assertEqual(prio['montant'], self._ligne(etat, 'CRITIQUE')['montant']
                         + self._ligne(etat, 'URGENT')['montant'])
        self.assertEqual(self._ligne(etat, 'ATTENTION_MOIS')['montant'], 10000)
        self.assertEqual(self._ligne(etat, 'ATTENTION_RETARD')['anterieur'], 50000)
        self.assertEqual(etat['total']['montant'],
                         prio['montant'] + self._ligne(etat, 'ATTENTION')['montant'])

    def test_les_montants_sont_ceux_de_la_liste_des_eleves(self):
        """Un seul calcul : l'alerte et la somme de chaque élève sont celles de sa fiche."""
        etat = self._etat()
        for l in etat['lignes']:
            for e in l.get('eleves', []):
                alerte = Eleve.objects.get(id=e['id']).situation_alerte(today=MARS)
                self.assertEqual(e['montant'], alerte['montant'])

    def test_telechargeable_en_pdf_et_en_excel(self):
        r = self.client.get('/api/eleves/etat-impayes/', {'export': 'pdf'})
        self.assertEqual((r.status_code, r['Content-Type']), (200, 'application/pdf'))
        texte = ''.join(p.extract_text() for p in PdfReader(BytesIO(r.content)).pages)
        self.assertIn('Sous-total prioritaires', texte)
        self.assertIn('TOTAL GÉNÉRAL', texte)

        r = self.client.get('/api/eleves/etat-impayes/', {'export': 'xlsx'})
        self.assertEqual(r.status_code, 200)
        from openpyxl import load_workbook
        ws = load_workbook(BytesIO(r.content)).active
        valeurs = [c for row in ws.iter_rows(values_only=True) for c in row if c]
        self.assertIn('TOTAL GÉNÉRAL', valeurs)
